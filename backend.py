# backend.py
import base64
import io
import json
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any

import cv2
import numpy as np
from PIL import Image
from filelock import FileLock
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from fastapi import FastAPI, File, UploadFile, Form, Request
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from ultralytics import YOLO

app = FastAPI()

# Static frontend
app.mount("/static", StaticFiles(directory="static"), name="static")


@app.get("/", response_class=HTMLResponse)
def index():
    return Path("static/index.html").read_text(encoding="utf-8")


# Storage
HISTORY_PATH = Path("history.jsonl")
LOCK_PATH = Path("history.jsonl.lock")
IMAGES_DIR = Path("history_images")
IMAGES_DIR.mkdir(exist_ok=True)


def append_history(record: Dict[str, Any]) -> None:
    LOCK_PATH.touch(exist_ok=True)
    with FileLock(str(LOCK_PATH)):
        with HISTORY_PATH.open("a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def read_history(limit: int = 50) -> List[Dict[str, Any]]:
    if not HISTORY_PATH.exists():
        return []
    LOCK_PATH.touch(exist_ok=True)

    with FileLock(str(LOCK_PATH)):
        lines = HISTORY_PATH.read_text(encoding="utf-8").splitlines()

    items: List[Dict[str, Any]] = []
    for line in reversed(lines):
        if not line.strip():
            continue
        try:
            items.append(json.loads(line))
        except Exception:
            continue
        if len(items) >= limit:
            break
    return items


def compute_summary(items: List[Dict[str, Any]]) -> Dict[str, Any]:
    total = len(items)
    with_unattended = sum(1 for x in items if int(x.get("unattended_count", 0)) > 0)
    total_unattended = sum(int(x.get("unattended_count", 0)) for x in items)

    bag_types: Dict[str, int] = {}
    for x in items:
        # суммируем по unattended (как раньше)
        for b in x.get("unattended", []):
            t = b.get("bag_type", "unknown")
            bag_types[t] = bag_types.get(t, 0) + 1

    top_types = sorted(bag_types.items(), key=lambda z: z[1], reverse=True)[:5]

    return {
        "total_requests": total,
        "requests_with_unattended": with_unattended,
        "total_unattended_found": total_unattended,
        "unattended_rate": (with_unattended / total) if total else 0.0,
        "top_bag_types": top_types,
    }


# Model
MODEL = YOLO(str(Path(__file__).with_name("person_baggage_best.pt")))

NAME2ID = {v: k for k, v in MODEL.names.items()}
PERSON_ID = NAME2ID.get("person", None)
BAG_ID = NAME2ID.get("baggage", None)

if PERSON_ID is None or BAG_ID is None:
    raise RuntimeError(f"Ожидаются классы 'person' и 'baggage', найдено: {MODEL.names}")

BAG_IDS = [BAG_ID]  # используем список как и раньше


# Geometry helpers
def box_center_xyxy(xyxy: np.ndarray):
    x1, y1, x2, y2 = xyxy
    return (x1 + x2) / 2.0, (y1 + y2) / 2.0


def box_area_xyxy(xyxy: np.ndarray):
    x1, y1, x2, y2 = xyxy
    return max(0.0, x2 - x1) * max(0.0, y2 - y1)


def norm_distance(bag_xyxy: np.ndarray, person_xyxy: np.ndarray) -> float:
    # distance / (sqrt(area_person) + sqrt(area_bag))
    bx, by = box_center_xyxy(bag_xyxy)
    px, py = box_center_xyxy(person_xyxy)
    d = float(((bx - px) ** 2 + (by - py) ** 2) ** 0.5)

    sp = float(box_area_xyxy(person_xyxy) ** 0.5)
    sb = float(box_area_xyxy(bag_xyxy) ** 0.5)
    denom = max(1e-6, sp + sb)
    return d / denom


def draw_unattended_only(image_rgb: np.ndarray, unattended: List[Dict[str, Any]]) -> np.ndarray:
    img_bgr = cv2.cvtColor(image_rgb, cv2.COLOR_RGB2BGR)
    for row in unattended:
        x1, y1, x2, y2 = map(int, [row["x1"], row["y1"], row["x2"], row["y2"]])
        label = f'{row["bag_type"]} {row["conf"]:.2f}'
        cv2.rectangle(img_bgr, (x1, y1), (x2, y2), (0, 0, 255), 2)
        cv2.putText(
            img_bgr,
            label,
            (x1, max(0, y1 - 6)),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.6,
            (0, 0, 255),
            2,
            cv2.LINE_AA,
        )
    return cv2.cvtColor(img_bgr, cv2.COLOR_BGR2RGB)


# ---------- API ----------
@app.post("/predict")
async def predict(
    request: Request,
    file: UploadFile = File(...),
    source: str = Form("upload"),       # upload | camera
    conf: float = Form(0.25),
    dist_thr: float = Form(1.0),
    person_conf: float = Form(0.5),
    min_area_ratio: float = Form(3.0),  # area_person / area_bag
):
    req_id = str(uuid.uuid4())
    ts = datetime.now().isoformat(timespec="seconds")

    content = await file.read()
    pil = Image.open(io.BytesIO(content)).convert("RGB")
    img_rgb = np.array(pil, dtype=np.uint8)

    classes_filter = [PERSON_ID] + BAG_IDS
    results = MODEL.predict(
        source=pil,
        conf=float(conf),
        imgsz=640,
        classes=classes_filter,
        verbose=False,
    )

    r0 = results[0]
    boxes = r0.boxes

    persons: List[Dict[str, Any]] = []
    bags: List[Dict[str, Any]] = []

    if boxes is not None and len(boxes) > 0:
        cls_ids = boxes.cls.cpu().numpy().astype(int)
        confs = boxes.conf.cpu().numpy()
        xyxy = boxes.xyxy.cpu().numpy()

        for i in range(len(cls_ids)):
            c = int(cls_ids[i])
            row = {
                "cls": c,
                "name": r0.names.get(c, str(c)),
                "conf": float(confs[i]),
                "x1": float(xyxy[i][0]),
                "y1": float(xyxy[i][1]),
                "x2": float(xyxy[i][2]),
                "y2": float(xyxy[i][3]),
            }

            if c == PERSON_ID and row["conf"] >= float(person_conf):
                persons.append(row)
            elif c in BAG_IDS:
                bags.append(row)

    person_xyxys = [
        np.array([p["x1"], p["y1"], p["x2"], p["y2"]], dtype=float) for p in persons
    ]

    unattended: List[Dict[str, Any]] = []

    for b in bags:
        b_xyxy = np.array([b["x1"], b["y1"], b["x2"], b["y2"]], dtype=float)
        bag_area = box_area_xyxy(b_xyxy)

        # фильтруем людей, которые слишком маленькие относительно багажа (скорее всего далеко)
        eligible_persons = []
        for pxy in person_xyxys:
            person_area = box_area_xyxy(pxy)
            if person_area / max(1e-6, bag_area) >= float(min_area_ratio):
                eligible_persons.append(pxy)

        if len(eligible_persons) == 0:
            min_nd = float("inf")
            is_unattended = True
        else:
            nds = [norm_distance(b_xyxy, pxy) for pxy in eligible_persons]
            min_nd = float(np.min(nds))
            is_unattended = (min_nd > float(dist_thr))

        if is_unattended:
            unattended.append(
                {
                    "bag_type": b["name"],
                    "conf": b["conf"],
                    "x1": b["x1"],
                    "y1": b["y1"],
                    "x2": b["x2"],
                    "y2": b["y2"],
                    "min_norm_dist": None if not np.isfinite(min_nd) else min_nd,
                }
            )

    # --- counts by type (all bags + unattended) ---
    bag_counts_by_type: Dict[str, int] = {}
    for b in bags:
        t = b["name"]
        bag_counts_by_type[t] = bag_counts_by_type.get(t, 0) + 1

    unattended_counts_by_type: Dict[str, int] = {}
    for u in unattended:
        t = u["bag_type"]
        unattended_counts_by_type[t] = unattended_counts_by_type.get(t, 0) + 1

    annotated_rgb = draw_unattended_only(img_rgb, unattended)

    # save annotated png to disk
    img_path = IMAGES_DIR / f"{req_id}.png"
    Image.fromarray(annotated_rgb).save(img_path, format="PNG")

    buf = io.BytesIO()
    Image.fromarray(annotated_rgb).save(buf, format="PNG")
    b64 = base64.b64encode(buf.getvalue()).decode("ascii")

    record = {
        "id": req_id,
        "ts": ts,
        "client_ip": request.client.host if request.client else None,
        "source": source,
        "filename": file.filename,  # фронт не показывает, но пусть хранится
        "params": {
            "conf": float(conf),
            "dist_thr": float(dist_thr),
            "person_conf": float(person_conf),
            "min_area_ratio": float(min_area_ratio),
        },
        "persons_count": len(persons),
        "bags_count": len(bags),
        "unattended_count": len(unattended),
        "bag_counts_by_type": bag_counts_by_type,
        "unattended_counts_by_type": unattended_counts_by_type,
        "unattended": unattended,
        "annotated_image": str(img_path),
    }
    append_history(record)

    return {
        "id": req_id,
        "annotated_png_base64": b64,
        "unattended": unattended,
        "counts": {
            "person": len(persons),
            "baggage": len(bags),
            "unattended": len(unattended),
        },
        "bag_counts_by_type": bag_counts_by_type,
        "unattended_counts_by_type": unattended_counts_by_type,
    }


@app.get("/history")
def history(limit: int = 50):
    items = read_history(limit=limit)
    summary = compute_summary(items)
    return {"summary": summary, "items": items}


@app.get("/history/image/{req_id}")
def history_image(req_id: str):
    p = IMAGES_DIR / f"{req_id}.png"
    if not p.exists():
        # вернём 404 без фейкового файла
        return FileResponse(str(p), status_code=404)
    return FileResponse(str(p), media_type="image/png")


@app.get("/report.xlsx")
def report_xlsx(limit: int = 1000):
    items = read_history(limit=limit)
    items = list(reversed(items))  # хронологически
    summary = compute_summary(items)

    wb = Workbook()
    ws = wb.active
    ws.title = "requests"

    headers = [
        "ts", "id", "source",
        "conf", "dist_thr", "person_conf", "min_area_ratio",
        "persons_count", "bags_count", "unattended_count",
        "baggage_counts_by_type",
        "unattended_counts_by_type",
    ]
    ws.append(headers)

    for r in items:
        params = r.get("params", {})
        bct = r.get("bag_counts_by_type") or {}
        uct = r.get("unattended_counts_by_type") or {}

        bct_str = ", ".join([f"{k}:{v}" for k, v in bct.items()]) if bct else "-"
        uct_str = ", ".join([f"{k}:{v}" for k, v in uct.items()]) if uct else "-"

        ws.append([
            r.get("ts"),
            r.get("id"),
            r.get("source"),
            params.get("conf"),
            params.get("dist_thr"),
            params.get("person_conf"),
            params.get("min_area_ratio"),
            r.get("persons_count"),
            r.get("bags_count"),
            r.get("unattended_count"),
            bct_str,
            uct_str,
        ])

    # автоширина колонок
    for col in range(1, len(headers) + 1):
        max_len = 0
        col_letter = get_column_letter(col)
        for cell in ws[col_letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(55, max(10, max_len + 2))

    # summary sheet
    ws2 = wb.create_sheet("summary")
    ws2.append(["total_requests", summary["total_requests"]])
    ws2.append(["requests_with_unattended", summary["requests_with_unattended"]])
    ws2.append(["total_unattended_found", summary["total_unattended_found"]])
    ws2.append(["unattended_rate", summary["unattended_rate"]])
    ws2.append([])
    ws2.append(["top_unattended_types", "count"])
    for t, c in summary["top_bag_types"]:
        ws2.append([t, c])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="report.xlsx"'},
    )
